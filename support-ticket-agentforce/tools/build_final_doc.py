#!/usr/bin/env python3
"""Build the submission document in the format required by the project guide.

Follows the 11-section structure of the supplied Template Format, with the two
leftover artifacts from the source project (vehicle/dealer testing paragraph and
the Apex/batch "what you'll learn" list) corrected, and real screenshots from
the org embedded.
"""
import pathlib
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

SHOTS = pathlib.Path("/home/killermachine/Desktop/study/nm/agentforce-shots")
OUT = pathlib.Path(
    "/home/killermachine/Desktop/study/nm/"
    "Customer Support Ticket Priority Prediction - Project Documentation.docx"
)
TITLE = ("Customer Support Ticket Priority Prediction and Automated "
         "Assignment System Using Agentforce")


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def p(doc, text):
    doc.add_paragraph(text)


def bullets(doc, items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def nums(doc, items):
    for i in items:
        doc.add_paragraph(i, style="List Number")


def table(doc, headers, rows, style="Light Grid Accent 1", font_size=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = style
    for i, x in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = x
        if font_size is not None:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = font_size
                    r.font.bold = True
    for row in rows:
        cells = t.add_row().cells
        is_total = (str(row[0]).lower() in ["total", ""] and "total" in str(row[1]).lower())
        for i, v in enumerate(row):
            cells[i].text = str(v)
            if font_size is not None:
                for p in cells[i].paragraphs:
                    for r in p.runs:
                        r.font.size = font_size
                        if is_total:
                            r.font.bold = True
    doc.add_paragraph()
    return t


def shot(doc, name, caption, width=6.2):
    path = SHOTS / name
    if not path.exists():
        r = doc.add_paragraph().add_run(f"[screenshot missing: {name}]")
        r.italic = True
        r.font.color.rgb = RGBColor(0x99, 0x33, 0x33)
        return
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def build():
    doc = Document()

    # ---- FIRST PAGE: COVER & TEAM DETAILS -----------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(TITLE.upper())
    r_title.bold = True
    r_title.font.size = Pt(17)
    r_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Salesforce-Based Support Ticket Intelligence and Conversational AI Automation")
    r_sub.italic = True
    r_sub.font.size = Pt(11.5)
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p_sub.paragraph_format.space_after = Pt(3)

    p_lbl = doc.add_paragraph()
    p_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_lbl = p_lbl.add_run("PROJECT DOCUMENTATION")
    r_lbl.bold = True
    r_lbl.font.size = Pt(11)
    p_lbl.paragraph_format.space_after = Pt(2)

    p_note = doc.add_paragraph()
    p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_note = p_note.add_run(
        "Detailed implementation, configuration, flow automation, Agentforce integration, "
        "testing, and team work distribution report."
    )
    r_note.font.size = Pt(9.5)
    r_note.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    p_note.paragraph_format.space_after = Pt(12)

    # College
    p_col = doc.add_paragraph()
    p_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_col_lbl = p_col.add_run("COLLEGE: ")
    r_col_lbl.bold = True
    r_col_lbl.font.size = Pt(12)
    r_col_name = p_col.add_run("GOVERNMENT COLLEGE OF ENGINEERING, ERODE")
    r_col_name.bold = True
    r_col_name.font.size = Pt(12)
    r_col_name.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    p_col.paragraph_format.space_after = Pt(12)

    # Team Members Block
    p_team = doc.add_paragraph()
    p_team.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_tl = p_team.add_run("TEAM LEADER: ")
    r_tl.bold = True
    r_tl.font.size = Pt(11)
    r_tln = p_team.add_run("PREMKUMAR P\n")
    r_tln.font.size = Pt(11)

    r_tm = p_team.add_run("TEAM MEMBERS:\n")
    r_tm.bold = True
    r_tm.font.size = Pt(11)
    for m in ["BHARATH KUMAR P", "VIKASH V", "PUGAZHENDHI P", "ASWIN M"]:
        p_team.add_run(f"\t\t\t\t• {m}\n")
    p_team.paragraph_format.space_after = Pt(10)

    # Team Work Distribution Summary Table
    p_summary_title = doc.add_paragraph("TEAM WORK DISTRIBUTION SUMMARY:")
    p_summary_title.runs[0].bold = True
    p_summary_title.runs[0].font.size = Pt(10.5)
    p_summary_title.paragraph_format.space_after = Pt(4)

    table(doc, ["S.No.", "Team Member", "Project Role", "Key Modules / User Stories", "Story Points"], [
        ("1", "Premkumar P (Team leader)", "AI & Backend Systems Lead", "Dev Setup, Agentforce Subagent & Action Mapping (USN-1, USN-6)", "11"),
        ("2", "Bharath Kumar P", "Data Architect & Admin", "Custom Object & Fields, Urgent Task Automation & Assignment (USN-2, USN-7)", "10"),
        ("3", "Vikash V", "Full-Stack Developer & QA Lead", "Relationships, UI Layouts & List Views, SLA Risk Logic (USN-3, USN-8)", "8"),
        ("4", "Pugazhendhi P", "Flow Automation Engineer", "Auto-Launched Flow Architecture, Ticket Retrieval, Messaging (USN-4, USN-9)", "10"),
        ("5", "Aswin M", "Business Logic & Testing Analyst", "Priority Decision Logic, Test Dataset & Functional Validation (USN-5)", "5"),
        ("Total", "5 Members", "Full Project Lifecycle Delivery", "9 Backlog User Stories, Flow & Agentforce Testing", "44"),
    ], font_size=Pt(9))

    doc.add_page_break()

    # ---- PAGE 2: USER STORY & INTRODUCTION ---------------------------------
    t = doc.add_heading(TITLE, level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    us = doc.add_paragraph()
    us.add_run("User Story: ").bold = True
    us.add_run(
        "The Customer Support Ticket Priority Prediction and Automated Assignment "
        "System uses Salesforce and Agentforce to analyze customer support ticket "
        "descriptions, automatically determine priority as High, Medium, or Low, and "
        "trigger backend automation for appropriate support handling. The solution uses "
        "the Support Ticket Intelligence custom object, an Auto-Launched Flow, and an "
        "Agentforce subagent named Support Ticket Priority Analysis. High-priority "
        "tickets create an urgent handling task and are assigned to a senior support "
        "agent, reducing manual effort and improving response time."
    )

    # ---- 1. INTRODUCTION ----------------------------------------------------
    h(doc, "1. Introduction")
    h(doc, "1.1 Project Overview", 2)
    p(doc, "Customer support teams receive a large number of tickets daily, while "
           "prioritization and assignment can be manual. This project provides a "
           "Salesforce-based solution that analyzes ticket information, classifies "
           "urgency, assigns the appropriate support level, and initiates automation "
           "for critical cases.")
    h(doc, "1.2 Purpose", 2)
    bullets(doc, [
        "Automatically classify support tickets as High, Medium, or Low priority.",
        "Retrieve the latest support ticket associated with a customer Account.",
        "Reduce manual prioritization and assignment effort.",
        "Create an urgent handling task for High-priority tickets.",
        "Use Agentforce to analyze ticket details and trigger backend automation.",
        "Improve ticket response and resolution efficiency.",
    ])

    # ---- 2. IDEATION --------------------------------------------------------
    h(doc, "2. Ideation Phase")
    h(doc, "2.1 Problem Statement", 2)
    table(doc, ["PS", "I am (Customer)", "I'm trying to", "But", "Because", "Which makes me feel"], [
        ("PS-1", "Support Agent", "focus on critical tickets first",
         "tickets are not always prioritized automatically",
         "ticket prioritization can be manual", "frustrated"),
        ("PS-2", "Customer", "get urgent issues resolved quickly",
         "critical issues may wait behind lower-priority tickets",
         "support teams handle large ticket volumes", "concerned"),
        ("PS-3", "Support Manager", "assign tickets to the right support level",
         "assignment may require manual decisions",
         "ticket urgency is not consistently analyzed", "overwhelmed"),
        ("PS-4", "Support Agent", "identify urgent issues from descriptions",
         "important keywords can be missed",
         "ticket descriptions contain different urgency signals", "uncertain"),
        ("PS-5", "Support Manager", "reduce SLA-related risk",
         "older unresolved tickets may remain unnoticed",
         "SLA risk needs additional checking", "concerned"),
    ])

    h(doc, "2.2 Empathy Map Canvas", 2)
    p(doc, "User Persona: Support Agent")
    bullets(doc, [
        "Think & Feel: Wants critical issues identified quickly so urgent customer "
        "problems receive attention.",
        "Hear: Receives customer issue descriptions and support requests.",
        "See: A queue of tickets of mixed urgency arriving throughout the day.",
        "Say & Do: Triages tickets, decides which to handle first, updates status.",
        "Pains: Manual triage is slow and inconsistent; urgent tickets can be missed.",
        "Gains: Automatic priority classification and assignment remove the guesswork.",
    ])

    h(doc, "2.3 Brainstorming & Idea Prioritization", 2)
    p(doc, "Ideas Considered")
    bullets(doc, ["Manual ticket registers", "Spreadsheet-based ticket tracking",
                  "Standalone support portal", "Salesforce CRM with Agentforce"])
    p(doc, "Final Selection — Salesforce CRM selected due to:")
    bullets(doc, [
        "Automation through Salesforce Flow",
        "AI-assisted ticket analysis through Agentforce",
        "Automatic priority classification",
        "Automatic urgent-task creation",
        "Reduced manual support effort",
    ])

    # ---- 3. REQUIREMENT ANALYSIS -------------------------------------------
    h(doc, "3. Requirement Analysis Phase")
    h(doc, "3.1 Customer Journey Map", 2)
    p(doc, "Stages: Ticket Creation → Account Identification → Ticket Retrieval → "
           "Description Analysis → Priority Classification → Agent Assignment → "
           "Urgent Task Creation → Support Handling → Resolution")

    h(doc, "3.2 Solution Requirements", 2)
    p(doc, "Functional Requirements")
    table(doc, ["FR No", "Requirement"], [
        ("FR-1", "Support ticket creation and management"),
        ("FR-2", "Account and Contact association"),
        ("FR-3", "Ticket description analysis"),
        ("FR-4", "Automatic priority classification"),
        ("FR-5", "High-priority task creation"),
        ("FR-6", "Support agent assignment"),
        ("FR-7", "SLA breach risk checking"),
        ("FR-8", "Agentforce conversational analysis"),
    ])
    p(doc, "Non-Functional Requirements")
    table(doc, ["NFR", "Description"], [
        ("Usability", "Clear Agentforce interaction and understandable priority messages"),
        ("Reliability", "Consistent priority classification based on configured conditions"),
        ("Performance", "Quick retrieval and analysis of the latest support ticket"),
        ("Security", "Controlled Salesforce access to customer and ticket information"),
        ("Maintainability", "Flow and Agentforce configuration can be updated"),
        ("Scalability", "Supports increasing support-ticket volumes through Salesforce automation"),
    ])

    h(doc, "3.3 Data Flow Diagram", 2)
    p(doc, "Flow: Customer / Support Request → Account → Latest Support Ticket → "
           "Ticket Description → Analyze Description → Priority Level → "
           "High/Medium/Low Decision → Agent Assignment → High-Priority Task "
           "(if applicable) → Final Response")
    p(doc, "Agentforce: User provides Account Name → Agentforce Support Ticket "
           "Priority Analysis → Auto-Launched Flow → Ticket Analysis → Priority and "
           "Assignment Output → Response to User.")

    h(doc, "3.4 Technology Stack", 2)
    table(doc, ["Layer", "Technology"], [
        ("Platform", "Salesforce"),
        ("Database", "Support_Ticket_Intelligence__c with Account, Contact, User and Task records"),
        ("Logic", "Auto-Launched Flow"),
        ("AI", "Agentforce"),
        ("Agent Configuration", "Support Ticket Priority Analysis subagent"),
        ("Automation", "Flow Decision, Assignment and Create Records"),
        ("Output", "Agentforce conversation response"),
    ])

    # ---- 4. DESIGN ----------------------------------------------------------
    h(doc, "4. Project Design Phase")
    h(doc, "4.1 Problem–Solution Fit", 2)
    table(doc, ["Problem", "Solution", "Benefit"], [
        ("Manual ticket prioritization", "Flow analyzes description and sets priority",
         "Faster identification of urgent cases"),
        ("Urgent issues may be missed", "High-priority keyword decision",
         "Critical cases receive attention"),
        ("Manual task creation", "Flow creates Task for High priority", "Reduced manual effort"),
        ("Manual support assignment", "Flow sets assigned support level",
         "More consistent handling"),
        ("SLA risk can be overlooked", "SLA Risk decision flags tickets older than two days",
         "Better risk awareness"),
        ("Need for conversational access", "Agentforce subagent", "AI-assisted interaction"),
    ])

    h(doc, "4.2 Proposed Solution", 2)
    table(doc, ["Parameter", "Description"], [
        ("Problem", "Manual ticket prioritization and assignment causing delays in "
                    "handling critical customer support issues"),
        ("Solution", "Salesforce-based Customer Support Ticket Priority Prediction and "
                     "Automated Assignment System"),
        ("Innovation", "Agentforce AI + Auto-Launched Flow for analyzing ticket "
                       "descriptions, automatically classifying priority, assigning "
                       "support level, and triggering urgent-task automation"),
        ("Impact", "Faster identification of critical tickets, reduced manual effort, "
                   "improved response time, and better support-team productivity"),
        ("Business Model", "Customer support operations involving customers, support "
                           "agents, support managers, and AI-assisted ticket handling"),
        ("Scalability", "Supports increasing support-ticket volumes through Salesforce "
                        "automation, with scope for additional assignment rules, SLA "
                        "monitoring, and priority indicators"),
    ])

    h(doc, "4.3 Solution Architecture", 2)
    p(doc, "Components: Users / Support Team → Agentforce → Support Ticket Priority "
           "Analysis → Auto-Launched Flow → Account & Ticket Data → Description "
           "Analysis → Priority Decision → Task / Assignment → Final Message → "
           "Agentforce Response")

    # ---- 5. PLANNING --------------------------------------------------------
    h(doc, "5. Project Planning & Scheduling")
    h(doc, "5.1 Project Planning & Work Breakdown Structure", 2)
    bullets(doc, [
        "Agile Scrum methodology with iterative sprint-based execution",
        "Epics decomposed into User Stories and estimated using Story Points",
        "Velocity tracking used for estimation and sprint capacity planning",
        "Work Package allocation across the 5-member engineering team",
    ])

    p(doc, "Work Breakdown Structure (WBS)")
    table(doc, ["Work Package", "Tasks / Activities", "Assigned Member", "Key Deliverable / Output"], [
        ("WP-01", "Developer Environment Configuration & SFDX Setup", "Premkumar P", "Salesforce Developer Org configured with Agentforce enabled"),
        ("WP-02", "Custom Object & Field Schema Modeling", "Bharath Kumar P", "Support_Ticket_Intelligence__c object with 11 custom fields"),
        ("WP-03", "Relationships, UI Layouts & Navigation", "Vikash V", "Account/Contact Lookups, Page Layout & All Tickets list view"),
        ("WP-04", "Auto-Launched Flow Architecture & Record Retrieval", "Pugazhendhi P", "Flow skeleton with Get Records elements for Account & Ticket"),
        ("WP-05", "Urgency Priority Decision Logic & Status Updates", "Aswin M", "Keyword decision rules (urgent/slow/etc.) & priority update back to record"),
        ("WP-06", "Urgent Task Creation & Senior Agent Assignment", "Bharath Kumar P", "High-priority Task creation & Senior Support Agent assignment"),
        ("WP-07", "SLA Breach Risk Calculation & Management", "Vikash V", "Decision rule for tickets > 2 days old & SLA_Breach_Risk__c flag"),
        ("WP-08", "Agentforce Subagent Creation & Flow Action Registration", "Premkumar P", "Support Ticket Priority Analysis subagent & active Flow action"),
        ("WP-09", "End-to-End Conversational Testing & Verification", "Pugazhendhi P & Aswin M", "Multi-branch agent testing, action messaging, and test evidence"),
    ], font_size=Pt(9))

    h(doc, "5.2 Team Work Distribution", 2)
    p(doc, "The project implementation and deliverables were systematically distributed among all 5 students based on technical competencies and project requirements. Each team member owned dedicated modules and backlog stories:")
    bullets(doc, [
        "Premkumar P (Team leader) — 11 Story Points: Architected the end-to-end solution and led developer setup. Configured the Salesforce environment, activated Agentforce Default, built and configured the 'Support Ticket Priority Analysis' subagent, defined conversational instructions, and registered the Auto-Launched Flow as an Agent Action.",
        "Bharath Kumar P — 10 Story Points: Led data architecture and task automation. Implemented the 'Support_Ticket_Intelligence__c' custom object schema and its 11 fields, configured field-level security permissions, built the automated High-priority Task creation element in Flow, and established Senior Support Agent routing.",
        "Vikash V — 8 Story Points: Designed data relationships linking tickets to customer Accounts and Contacts, customized page layouts, created the 'All Tickets' list view for full visibility, implemented the SLA breach risk decision logic for tickets older than 2 days, and ensured metadata traceability.",
        "Pugazhendhi P — 10 Story Points: Developed the Auto-Launched Flow architecture ('Support_Ticket_Intelligence'), constructed Get Records elements for Account identification and latest ticket retrieval, defined Flow output variables, mapped conversational action messages, and conducted conversational verification.",
        "Aswin M — 5 Story Points: Engineered the keyword-based priority classification decision logic (High, Medium, Low branches), formulated keyword criteria ('urgent', 'not working', 'failure', 'slow', 'delay'), created multi-account sample datasets (Acme Corporation, Globex Systems, Initech Solutions), and validated priority assignment accuracy.",
    ])

    p(doc, "Team Work Distribution Summary")
    table(doc, ["S.No.", "Team Member", "Project Role", "User Stories Handled", "Story Points"], [
        ("1", "Premkumar P (Team leader)", "AI & Backend Systems Lead", "USN-1, USN-6", "11"),
        ("2", "Bharath Kumar P", "Data Architect & Admin", "USN-2, USN-7", "10"),
        ("3", "Vikash V", "Full-Stack Developer & QA Lead", "USN-3, USN-8", "8"),
        ("4", "Pugazhendhi P", "Flow Automation Engineer", "USN-4, USN-9", "10"),
        ("5", "Aswin M", "Business Logic & Testing Analyst", "USN-5", "5"),
        ("Total", "5 Team Members", "Full Application Lifecycle", "9 Backlog Stories", "44"),
    ], font_size=Pt(9))

    h(doc, "5.3 Product Backlog, Sprint Schedule, and Estimation", 2)
    table(doc, ["Sprint", "Functional Requirement (Epic)", "User Story No.",
                "User Story / Task", "Story Points", "Priority", "Team Member"], [
        ("Sprint-1", "Developer Setup", "USN-1",
         "As a system administrator, I want to configure the Salesforce environment for "
         "Support Ticket Intelligence so that the support automation solution can be developed.",
         "3", "High", "Premkumar P"),
        ("Sprint-2", "Data Modeling", "USN-2",
         "As an admin, I want to create the Support Ticket Intelligence object with required "
         "fields and relationships so that customer support ticket information is stored systematically.",
         "5", "High", "Bharath Kumar P"),
        ("Sprint-2", "Data Modeling", "USN-3",
         "As a support user, I want support tickets to be associated with Accounts and "
         "Contacts so that the latest customer ticket can be identified for analysis.",
         "3", "High", "Vikash V"),
        ("Sprint-3", "Automation", "USN-4",
         "As a support agent, I want an Auto-Launched Flow to retrieve the latest ticket and "
         "analyze its description so that ticket priority can be determined automatically.",
         "5", "High", "Pugazhendhi P"),
        ("Sprint-3", "Automation", "USN-5",
         "As a support agent, I want tickets to be automatically classified as High, Medium, "
         "or Low based on configured urgency conditions so that critical issues receive attention first.",
         "5", "High", "Aswin M"),
        ("Sprint-4", "Agentforce & Automation", "USN-6",
         "As an AI Agent (Agentforce), I want to analyze support ticket details and trigger the "
         "backend Flow so that manual ticket prioritization is reduced.",
         "8", "High", "Premkumar P"),
        ("Sprint-4", "Task & Assignment", "USN-7",
         "As a support manager, I want High-priority tickets to automatically create an urgent "
         "handling task and assign the appropriate support level so that critical tickets are handled quickly.",
         "5", "High", "Bharath Kumar P"),
        ("Sprint-5", "SLA Management", "USN-8",
         "As a support manager, I want the system to check SLA breach risk for older unresolved "
         "tickets so that potential SLA issues can be identified.",
         "5", "Medium", "Vikash V"),
        ("Sprint-6", "Agentforce & Conversational Support", "USN-9",
         "As a support user, I want to provide an Account Name to Agentforce and receive the "
         "ticket ID, priority, assigned agent, and action message so that I can understand the "
         "ticket handling outcome conversationally.",
         "5", "Medium", "Pugazhendhi P"),
    ], font_size=Pt(9))

    h(doc, "5.4 Project Tracker, Velocity & Burndown Chart", 2)
    table(doc, ["Sprint", "Total Story Points", "Duration", "Sprint Start Date",
                "Sprint End Date (Planned)", "Story Points Completed"], [
        ("Sprint-1", "3", "4 Days", "14 Sep 2026", "17 Sep 2026", "3"),
        ("Sprint-2", "8", "5 Days", "18 Sep 2026", "22 Sep 2026", "8"),
        ("Sprint-3", "10", "5 Days", "23 Sep 2026", "27 Sep 2026", "10"),
        ("Sprint-4", "13", "3 Days", "28 Sep 2026", "30 Sep 2026", "13"),
        ("Sprint-5", "5", "1 Day", "01 Oct 2026", "01 Oct 2026", "5"),
        ("Sprint-6", "5", "2 Days", "02 Oct 2026", "03 Oct 2026", "5"),
    ], font_size=Pt(9))
    doc.add_paragraph()

    doc.add_page_break()

    # ---- 6. DEVELOPMENT -----------------------------------------------------
    h(doc, "6. Project Development Phase")

    h(doc, "What you'll learn", 2)
    # NOTE: the supplied template listed Apex Triggers / Batch Apex / Scheduled
    # Apex here, which belong to the source project and are not used in this one.
    nums(doc, ["Data Modelling", "Fields and Relationships",
               "Auto-Launched Flows", "Flow Decisions and Record Creation",
               "Agentforce Subagents and Flow Actions",
               "Testing an Agentforce automation end to end"])

    h(doc, "Milestone 1: Salesforce Account", 2)
    p(doc, "Activity 1: Creating Developer Account")
    p(doc, "Creating a developer org in Salesforce.")
    nums(doc, [
        "Go to https://developer.salesforce.com/signup",
        "On the sign-up form enter First name, Last name, Email, Role: Developer, "
        "Company: College Name, Country: India, Postal Code, and a Username in the "
        "format username@organization.com",
        "Click Sign me up.",
    ])
    p(doc, "Activity 2: Account Activation")
    nums(doc, [
        "Open the inbox of the email used at sign-up and click Verify Account. "
        "The email may take 5–10 minutes.",
        "Set a password and answer a security question, then click Change Password.",
        "You are redirected to the Salesforce setup page.",
    ])
    p(doc, "Note: the project requires an org with Agentforce available. Confirm this "
           "before starting — the Agentforce milestone cannot be completed without it.")
    shot(doc, "10_object_manager_list.png",
         "Setup → Object Manager, the starting point for the data model.")

    h(doc, "Milestone 2: Data Management — Objects and Fields", 2)
    p(doc, "Custom Object")
    p(doc, "Object Name: Support Ticket Intelligence")
    p(doc, "API Name: Support_Ticket_Intelligence__c")
    doc.add_paragraph()
    p(doc, "Fields")
    table(doc, ["Field Label", "API Name", "Data Type", "Description"], [
        ("Ticket Number", "Ticket_Number__c", "Auto Number", "TKT-{0000}"),
        ("Customer", "Customer__c", "Lookup (Account)", "Related Account"),
        ("Contact", "Contact__c", "Lookup (Contact)", "Customer contact"),
        ("Issue Type", "Issue_Type__c", "Picklist", "Technical, Billing, General"),
        ("Description", "Description__c", "Long Text Area", "Issue details"),
        ("Priority Level", "Priority_Level__c", "Picklist", "Low, Medium, High"),
        ("Status", "Status__c", "Picklist", "New, In Progress, Resolved"),
        ("Created Date", "Created_Date__c", "Date", "Ticket date"),
        ("Assigned To", "Assigned_To__c", "Lookup (User)", "Support agent"),
        ("SLA Breach Risk", "SLA_Breach_Risk__c", "Checkbox", "Risk flag"),
        ("Resolution Time (hrs)", "Resolution_Time__c", "Number", "Time taken"),
    ])
    shot(doc, "11_object_detail.png", "The completed object in Object Manager.")
    shot(doc, "01_objectmanager_fields.png",
         "Fields & Relationships showing the completed field set.")

    h(doc, "Note: field-level security", 2)
    p(doc, "Fields created through the Metadata API are not automatically visible to "
           "any profile — not even System Administrator. Until field-level security is "
           "granted, querying the field fails with \"No such column\", which reads like "
           "a missing field rather than a permissions problem. Create a permission set "
           "covering the object and all its fields, then assign it to your user before "
           "building the Flow.")
    p(doc, "Create a list view as well. Without one, the object's tab renders empty "
           "even when records exist.")

    h(doc, "Milestone 3: Sample Data", 2)
    p(doc, "Create records before building the Flow so it has something to analyse.")
    table(doc, ["Account", "Ticket", "Description"], [
        ("Acme Corporation", "Production outage",
         "Urgent issue - production system is not working and all orders are failing."),
        ("Globex Systems", "Order processing lag",
         "The system is very slow and there is a delay in processing orders."),
        ("Initech Solutions", "Roadmap question",
         "Just a general query about the product roadmap."),
    ])
    shot(doc, "14_ticket_list_all.png",
         "All Tickets list view. TKT-0001 shows Priority High with SLA Breach Risk "
         "ticked; the newer tickets sit at Medium and Low.")
    shot(doc, "12_ticket_record.png",
         "A ticket record. The Urgent Ticket Handling tasks in the Activity panel "
         "were created by the Flow, not by hand.")

    h(doc, "Milestone 4: Auto-Launched Flow", 2)
    p(doc, "The Flow type must be Auto-Launched (no trigger). Agentforce can only "
           "invoke an Auto-Launched Flow, so a record-triggered Flow would not work here.")
    p(doc, "Variables")
    table(doc, ["Variable", "Direction", "Data Type", "Description"], [
        ("varAccountName", "Input", "Text", "Account name supplied by the user"),
        ("varAccountId", "Output", "Text", "Resolved account Id"),
        ("varTicketId", "Output", "Text", "Latest ticket Id"),
        ("varPriorityLevel", "Output", "Text", "High / Medium / Low"),
        ("varAssignedTo", "Output", "Text", "Agent assigned to the ticket"),
        ("varActionMessage", "Output", "Text", "Message returned in conversation"),
    ])
    p(doc, "Step-by-step implementation")
    nums(doc, [
        "Get Account — Get Records on Account where Name equals varAccountName. "
        "Select Only the first record.",
        "Store Account Id — Assignment setting varAccountId from Get_Account.Id.",
        "Get Ticket — Get Records on Support Ticket Intelligence where Customer equals "
        "varAccountId, sorted by Created Date descending, Only the first record.",
        "Store Ticket Id — Assignment setting varTicketId from Get_Ticket.Id.",
        "Check SLA Risk — Decision. If Created_Date__c is more than two days old, "
        "update the ticket to set SLA_Breach_Risk__c.",
        "Analyze Description — Decision with three outcomes: High Priority if the "
        "description contains urgent, not working or failure; Medium Priority if it "
        "contains issue, slow or delay; the default outcome is Low Priority.",
        "Set Priority Level — Assignment setting varPriorityLevel on each branch.",
        "Update Priority — Update Records writing Priority_Level__c back to the ticket.",
        "Is High Priority — Decision on varPriorityLevel.",
        "Create Task — only on the High branch. Subject Urgent Ticket Handling, "
        "Priority High, Status Not Started, related to the ticket Id.",
        "Set Assigned Agent — Assignment setting varAssignedTo to Senior Support Agent.",
        "Set Final Message — Assignment setting varActionMessage per priority.",
    ])
    p(doc, "Important: every Get Records element must have Only the first record "
           "selected. Without it the element returns a collection, and referencing "
           "Get_Ticket.Description__c then fails with a misleading error saying the "
           "field does not exist.")
    p(doc, "Final Flow Structure")
    p(doc, "Start → Get Account → Assignment (Account Id) → Get Ticket → Assignment "
           "(Ticket Id) → Check SLA Risk → Analyze Description → Assignment (Priority "
           "Level) → Update Priority → Is High Priority → Create Task → Assignment "
           "(Assigned Agent) → Set Final Message → End")
    shot(doc, "05_flow_builder_active.png",
         "The completed Flow, active, with all elements wired.")
    p(doc, "Final Outcome")
    bullets(doc, [
        "Takes an Account Name as input",
        "Fetches the latest support ticket for that account",
        "Analyses the description using keywords",
        "Sets the priority automatically and writes it back to the record",
        "Creates a task for urgent cases only",
        "Flags SLA breach risk for tickets older than two days",
        "Returns the result to Agentforce",
    ])

    h(doc, "Milestone 5: Agentforce Configuration", 2)
    nums(doc, [
        "Go to the Gear icon and click Salesforce Go.",
        "Search Agentforce Default, select it, click Get Started → Turn On → Confirm.",
        "Search Agents in the Quick Find box and open the agent, then click Open in "
        "Builder.",
        "Deactivate the agent — agent changes are only permitted while deactivated.",
        "In the Subagents panel choose New → New Subagent. Describe the job in plain "
        "language and the Builder generates the name, classification description, "
        "scope and instructions.",
        "Set the Name to Support Ticket Priority Analysis and the API name to "
        "Support_Ticket_Priority_Analysis.",
        "Finish the subagent, then open This Subagent's Actions → New → Create New Action.",
        "Reference Action Type: Flow. Reference Action: Support Ticket Intelligence.",
        "Set the Loading Text; on the input varAccountName enable Collect data from "
        "user; on all five outputs enable Show in conversation.",
        "Click Finish, then Activate the agent. Choose Ignore & Activate if the "
        "configuration checklist warns about Data Cloud or channels.",
    ])
    shot(doc, "06_agentforce_agents.png",
         "Agentforce Agents in Setup, with the Agentforce toggle switched on.")
    shot(doc, "20_agent_active_subagent.png",
         "Support Ticket Priority Analysis configured on the agent, with the agent active.")
    shot(doc, "21_subagent_actions_flow.png",
         "The subagent's actions list showing the Support Ticket Intelligence flow attached.")
    shot(doc, "22_subagents_list.png",
         "The subagents on the agent, including the new one.")

    h(doc, "Note: registering a flow as an agent action", 2)
    p(doc, "A Flow does not appear in the subagent's action picker until it has been "
           "registered as an agent action. On the subagent creation wizard the action "
           "list shows \"0 items\" even for a valid, active flow — including when "
           "filtered by name. This is not a fault in the Flow.")
    p(doc, "The registration step is New → Create New Action on the This Subagent's "
           "Actions tab. That dialog accepts Reference Action Type = Flow, and the "
           "flow appears there. Once created it shows in the subagent's action list.")

    doc.add_page_break()

    # ---- 7. TESTING (corrected) --------------------------------------------
    h(doc, "7. Functional & Performance Testing")
    p(doc, "During testing, screenshots were captured from Salesforce to validate the "
           "Support Ticket Intelligence object, the Auto-Launched Flow, ticket "
           "priority classification, SLA breach flagging, High-priority task creation, "
           "agent assignment, and the Agentforce subagent conversation. All features "
           "were verified to ensure correct functionality and performance.")
    doc.add_paragraph()
    table(doc, ["Account", "Description keywords", "Priority", "Task created", "SLA risk"], [
        ("Acme Corporation", "urgent, not working", "High", "Yes", "True"),
        ("Globex Systems", "slow, delay", "Medium", "No", "False"),
        ("Initech Solutions", "(none)", "Low", "No", "False"),
    ])
    p(doc, "Only the High-priority ticket produced a task, which is the specified "
           "behaviour. The Flow also writes the predicted priority back to the ticket "
           "record and flags SLA breach risk for tickets raised more than two days ago.")

    h(doc, "End-to-end test through the agent", 2)
    p(doc, "The Flow was also exercised through the agent conversation, which is the "
           "real test of the system: the agent must pick the right subagent, call the "
           "Flow with the user's input, and report the result in plain language.")
    p(doc, "Asking \"What is the priority of the support ticket for Acme Corporation?\" "
           "produced: Subagent Selected — Support Ticket Priority Analysis, then Action "
           "Launched — Support Ticket Intelligence (0.63 sec) with input "
           "{\"varAccountName\": \"Acme Corporation\"} and output returning "
           "varPriorityLevel \"High\".")
    shot(doc, "23_agent_test_high.png",
         "Agent conversation: the High branch, showing subagent and action selection.")
    shot(doc, "24_agent_test_low.png",
         "Agent conversation: the Low branch for Initech Solutions.")

    # ---- 8. ADVANTAGES ------------------------------------------------------
    h(doc, "8. Advantages & Disadvantages")
    p(doc, "Advantages")
    bullets(doc, [
        "Automatic Ticket Prioritization — classifies support tickets as High, Medium, "
        "or Low based on ticket description.",
        "Faster Identification of Critical Issues — helps support agents focus on "
        "urgent tickets first.",
        "Reduced Manual Effort — automates ticket analysis, priority setting, task "
        "creation, and assignment.",
        "Intelligent Ticket Handling — Agentforce analyzes ticket details and initiates "
        "the appropriate backend automation.",
        "Automatic Urgent Task Creation — High-priority tickets automatically create "
        "an Urgent Ticket Handling task.",
        "Improved Support Assignment — assigns High-priority tickets to the appropriate "
        "support level, including a senior support agent.",
        "Conversational Support — support users can interact with Agentforce to analyze "
        "ticket information and receive the outcome.",
        "Better Response Time — automation helps reduce delays in handling critical "
        "customer issues.",
        "SLA Risk Awareness — the solution checks older unresolved tickets for "
        "potential SLA breach risk.",
        "Scalable Support Operations — Salesforce automation can support increasing "
        "volumes of support tickets.",
        "Improved Team Productivity — reduces repetitive work and allows support teams "
        "to concentrate on ticket resolution.",
    ])
    p(doc, "Limitations")
    bullets(doc, [
        "Keyword-Based Priority Logic — the documented priority classification depends "
        "on configured keywords such as urgent, not working, failure, issue, slow, and delay.",
        "Limited Understanding of Context — the current implementation may not capture "
        "every possible urgency signal because the logic is based on configured conditions.",
        "Dependence on Accurate Ticket Data — correct Account, Contact, and ticket "
        "information is required for successful analysis.",
        "Flow Configuration Dependency — changes to ticket handling require maintenance "
        "of the Auto-Launched Flow and its logic.",
        "Agentforce Configuration Dependency — responses and actions depend on properly "
        "configured subagent instructions and Flow actions.",
        "Limited Scope of Agentforce — the subagent is focused on ticket analysis, "
        "priority determination, and backend automation; unrelated requests such as "
        "billing or subscription management are outside its scope.",
        "SLA Monitoring is Limited — the implementation includes an SLA-risk check "
        "rather than comprehensive SLA monitoring and escalation.",
        "Assignment Logic May Need Expansion — additional support teams or more "
        "sophisticated assignment rules would require further configuration.",
        "Maintenance Required — priority keywords, Flow variables, Agentforce outputs, "
        "and assignment logic may need updating as support requirements evolve.",
        "Limited Analytics in Current Scope — advanced analytics for priority trends and "
        "resolution performance are identified as future scope.",
    ])

    # ---- 9. CONCLUSION ------------------------------------------------------
    h(doc, "9. Conclusion")
    p(doc, "The Customer Support Ticket Priority Prediction and Automated Assignment "
           "System demonstrates how Salesforce Flow and Agentforce can work together "
           "to reduce manual support-ticket handling. The solution retrieves the latest "
           "ticket for an Account, analyzes its description, classifies priority, "
           "creates an urgent task for High-priority cases, assigns a support level, "
           "and returns a clear conversational response. This supports faster ticket "
           "handling, reduced manual effort, intelligent prioritization, and better "
           "team productivity.")

    # ---- 10. FUTURE SCOPE ---------------------------------------------------
    h(doc, "10. Future Scope")
    bullets(doc, [
        "Add more priority indicators and business rules.",
        "Extend Agentforce outputs with additional ticket information.",
        "Add more sophisticated SLA monitoring and escalation.",
        "Support additional support teams or assignment rules.",
        "Introduce analytics for priority trends and resolution performance.",
        "Add more variables and configure them in the Agentforce action for richer outputs.",
    ])

    # ---- 11. APPENDIX -------------------------------------------------------
    h(doc, "11. Appendix")
    bullets(doc, [
        "A. Custom Object: Support Ticket Intelligence",
        "B. Key Salesforce Records: Account, Contact, User, Task",
        "C. Flow: Support_Ticket_Intelligence — Auto-Launched Flow",
        "D. Agentforce Subagent: Support Ticket Priority Analysis",
        "E. Agentforce Action: Flow-based Support Ticket Intelligence action",
    ])

    doc.add_paragraph()
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = end.add_run("Thank You")
    r.bold = True
    r.font.size = Pt(14)

    return doc


def main():
    doc = build()
    doc.save(str(OUT))
    print("written:", OUT)
    print("size   :", f"{OUT.stat().st_size/1024:.1f} KB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
